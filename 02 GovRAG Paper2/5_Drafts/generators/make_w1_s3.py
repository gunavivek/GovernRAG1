#!/usr/bin/env python3
# W1 — §III The GovernRAG System: review draft for Vivek.
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/tmp/gr/W1_SectionIII_GovernRAG_System_DRAFT.docx"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
INK  = RGBColor(0x1A, 0x1A, 0x1A)
MUT  = RGBColor(0x55, 0x55, 0x55)
GRN  = RGBColor(0x1E, 0x6B, 0x34)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)

def P(runs, size=11, before=0, after=8, align=None, indent=None):
    if isinstance(runs, str): runs = [(runs, {})]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    for text, opts in runs:
        r = p.add_run(text)
        r.font.name = "Calibri"; r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", False); r.font.italic = opts.get("italic", False)
        r.font.color.rgb = opts.get("color", INK)
    return p

P([("W1 — §III draft v2 for review · 25 July 2026 · §III.A rulings applied "
    "(D amended; M = modelling pipeline; enterprise build-once rationale added; "
    "B-series excluded) · not yet poured into the IEEE shell",
    {"italic": True, "size": 9, "color": MUT})], after=4)
P([("Voice: fresh-written, no Paper-1 text; British spellings per the locked guide; "
    "every technical statement traceable to spine v2.0 §5 / Governance_Spectrum_Design.",
    {"italic": True, "size": 9, "color": MUT})], after=14)

P([("III. THE GOVERNRAG SYSTEM", {"bold": True, "size": 13, "color": NAVY})], after=8)

P([("A. Design position", {"bold": True, "italic": True})], after=4)

P("GovernRAG rests on a single architectural commitment: the decision of which "
  "retrieved evidence a language model may condition on is made by a dedicated "
  "enforcement component, not by retrieval ranking and not by the model itself. The "
  "component is a Reference Monitor in the operating-systems sense — it mediates "
  "every access of the generator to evidence, it cannot be bypassed while governance "
  "is active, and it logs every decision it takes. What distinguishes GovernRAG's "
  "monitor is its origin: its authorisation rules are derived from a reference "
  "ontology, so that what counts as admissible evidence is defined by an explicit "
  "domain vocabulary rather than by similarity scores.")

P("The system comprises four pipelines (Fig. 1). A reference ontology (R) supplies "
  "the domain vocabulary. A governance-definition pipeline (D) compiles governance "
  "policy together with domain-specific data into a per-record Governance Manifest. "
  "A modelling pipeline (M) constructs a governed index from the corpus. A query "
  "pipeline (Q) serves questions under a selected governance level. The build/serve "
  "split mirrors how documents are used in an enterprise — a document is written "
  "once and queried many times, not written to answer a single question — so "
  "governance work that depends only on the corpus is paid once at build time, and "
  "each query pays only the marginal cost of authorisation and synthesis; Section "
  "VI quantifies this economy.")

P([("B. The Governance Manifest", {"bold": True, "italic": True})], before=4, after=4)

P("The D pipeline terminates in the Governance Manifest (D5), a machine-readable "
  "policy artefact produced per record. The manifest carries three elements: the "
  "authorised predicates for the record, the ontology rules in force, and a domain-"
  "affinity specification. It is the single point of policy definition — the build "
  "pipeline consumes it when constructing the governed index, and the Reference "
  "Monitor consumes it when authorising evidence at serve time — so a policy change "
  "is made in one place and enforced in both. In our evaluation the reference "
  "ontology is instantiated with an enterprise reference vocabulary of 553 concepts "
  "derived from the Business Architecture Body of Knowledge (BIZBOK); the "
  "architecture itself is ontology-agnostic, and Section IV treats the "
  "instantiation as an experimental condition.")

P([("C. Build and query pipelines", {"bold": True, "italic": True})], before=4, after=4)

P("The modelling pipeline (M1–M5) performs governed chunking, concept and triple "
  "extraction, concept-graph construction with recursive bridge discovery, "
  "alignment of extracted concepts against the reference ontology, and assembly of "
  "the governed index. Alignment at M4 is the step that later gives governance its "
  "selectivity: each extracted concept is either matched to the ontology or "
  "recorded as unmatched, and this alignment status travels with the evidence into "
  "the index.")

P("At serve time the query pipeline runs intent gating, signature extraction, "
  "governed graph retrieval, and residual retrieval, producing candidate evidence "
  "for the question. The Q5 Reference Monitor then examines each candidate under "
  "the manifest and the active governance level, admitting or dropping it and "
  "writing one log entry per decision — the evidence item, the rule applied, and "
  "the outcome. Q6 cited synthesis enforces the level's answer gate on the admitted "
  "evidence: the model answers with citations drawn only from admitted evidence, or "
  "the system refuses. The gate is fail-closed — when admitted evidence does not "
  "satisfy the active level's requirement, refusal is the designed outcome, not an "
  "error state. The two stages separate concerns that single-gate designs merge: "
  "Q5 decides what may be used; Q6 decides whether what remains suffices to answer. "
  "From the logs alone, an auditor can reconstruct which evidence was admitted, "
  "which was dropped and under which rule, and which citations support each "
  "answer, without re-running the system.")

P([("D. Governance levels and axis variants", {"bold": True, "italic": True})],
  before=4, after=4)

P("Strictness is set by a governance level, chosen per deployment as a matter of "
  "policy. Five levels are defined, and they nest: each level retains every "
  "constraint of the level below and adds one. G0 (naive) bypasses the query "
  "pipeline entirely — retrieval passes straight to the model — and serves as the "
  "ungoverned baseline; because G0 differs from the governed levels only in the "
  "bypass, measured differences are attributable to governance rather than to the "
  "serving model or retrieval. G1 (Cited) requires at least one retrieved chunk and "
  "a citation in the answer. G2 (Grounded) requires at least one authorised chunk — "
  "a chunk that passes the monitor's predicate filter — together with the citation "
  "requirement. G3 (Strict) requires at least one authorised symbolic triplet. G4 "
  "(Corroborated) requires multiple authorised triplets, or one triplet supported "
  "by an independent agreeing authorised chunk. Under this construction the "
  "refusal behaviour of the system follows the declared policy: turning the dial "
  "from G1 to G4 only ever tightens what evidence may be used and what an answer "
  "must rest on.")

P("Within a level, two ontology-derived axis variants filter the admitted evidence "
  "further: ontology conformance restricts triplets by their alignment status "
  "against the reference ontology, and domain affinity restricts chunks by a "
  "domain-affinity score. The levels and the two axes together define the "
  "configuration space evaluated in this study — nine configurations per question "
  "(Section IV). Whether the axes can be exercised at all proved to be a property "
  "of the corpus, a finding reported in Section V.")

P([("[Figure 1 near §III.A — architecture, exported from "
    "3_Figures\\GovernRAG_Design_Diagram2b.pptx: R + D pipeline above the "
    "build-time M pipeline, D5 feeding M1 and Q5, serve-time Q pipeline with the "
    "Q5 monitor and Q6 gate.]", {"italic": True, "size": 10, "color": GRN})],
  before=6, after=14)

P([("Reviewer notes (not part of the manuscript)", {"bold": True, "size": 11,
    "color": NAVY})], after=4)
for t in [
 "Length: ~870 words ≈ 1.45 pp two-column with Figure 1 — inside the 1.5 pp budget.",
 "Every level definition, pipeline name, and the 553-concept instantiation checked against spine v2.0 §5 verbatim semantics; no Paper-1 sentences reused.",
 "Term discipline: Reference Monitor, Governance Manifest, governance level, admitted evidence, authorised chunk — one term each, throughout.",
 "The empty-evidence refusal state is stated here only as the fail-closed principle; its operational definition (questions whose admitted evidence yields no triplets) stays in §IV to avoid duplication.",
 "Open item for the final pass: the shell's abstract uses American spellings (authorization, behavior) while the locked voice guide mandates British (authorised, behaviour). One convention must win at W6 — flagging now, no action taken.",
 "Figure 1 export from your edited Diagram2b master pptx is pending — I will render it at pour time.",
]:
    P([("– ", {}), (t, {"size": 10})], after=3)

doc.save(OUT)
print("saved", OUT)
